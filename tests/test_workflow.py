from pathlib import Path
import yaml
from docx import Document
from PIL import Image
from gipa.workflow import approve, plan, prompts, validate_project

ROOT=Path(__file__).resolve().parents[1]

def test_end_to_end_prompt_workflow(tmp_path: Path) -> None:
    out=tmp_path/".ai-assets"
    state=plan(ROOT/"examples"/"SAMPLE_GDD.md",out)
    assert state["status"]=="awaiting_human_approval"
    manifest=yaml.safe_load((out/"ASSET_MANIFEST.yaml").read_text(encoding="utf-8"))
    assert manifest["assets"]
    assert any(a["category"]=="environment_texture" for a in manifest["assets"])
    assert any(a["category"]=="environment_background" for a in manifest["assets"])
    approve(out)
    prompts(out)
    assert (out/"prompts"/"BATCH_PROMPTS.md").exists()
    assert (out/"prompts"/"CHATGPT_BATCH_PROMPTS.md").exists()
    assert (out/"prompts"/"GEMINI_BATCH_PROMPTS.md").exists()
    assert validate_project(out)==[]

def test_missing_art_style_blocks_plan(tmp_path: Path) -> None:
    gdd=tmp_path/"plain.md"
    gdd.write_text("# Plain Game\n\n## Rules\nMove pieces to reach the exit.",encoding="utf-8")
    assert plan(gdd,tmp_path/"out")["status"]=="blocked"

def test_explicit_dimensions_and_alpha_override_fallback(tmp_path: Path) -> None:
    gdd=tmp_path/"dimensions.md"
    gdd.write_text(
        "# Dimension Game\n\n## Art Direction\nStylized 3D with clean shapes and cyan lighting.\n\n"
        "## Assets\nCreate a floor tile at 1536 x 768 with an opaque background.\n"
        "Create a hero character at 1024 × 2048 with a transparent background.\n",
        encoding="utf-8",
    )
    out=tmp_path/"out"
    plan(gdd,out)
    manifest=yaml.safe_load((out/"ASSET_MANIFEST.yaml").read_text(encoding="utf-8"))
    floor=next(a for a in manifest["assets"] if a["concept"]=="floor_tile")
    hero=next(a for a in manifest["assets"] if a["concept"]=="player_character")
    assert (floor["dimensions"]["width"],floor["dimensions"]["height"])==(1536,768)
    assert floor["dimensions"]["source"]["type"]=="gdd"
    assert floor["output"]["alpha"] is False
    assert (hero["dimensions"]["width"],hero["dimensions"]["height"])==(1024,2048)
    assert hero["output"]["alpha"] is True

def test_docx_title_and_assets_are_supported(tmp_path: Path) -> None:
    gdd=tmp_path/"02_Tiny_Heist_Game_Concept.docx"
    doc=Document()
    doc.add_paragraph("Tiny Heist",style="Title")
    doc.add_heading("1. Purpose and Locked Constraints",level=1)
    doc.add_heading("10. Art and Audio Direction",level=1)
    doc.add_paragraph("Premium stylized 3D with strong top-down readability and controlled cyan lighting.")
    doc.add_heading("Assets",level=1)
    doc.add_paragraph("The thief avoids guards, security cameras, lasers and doors on a seamless floor tile.")
    doc.save(gdd)
    out=tmp_path/"out"
    plan(gdd,out)
    manifest=yaml.safe_load((out/"ASSET_MANIFEST.yaml").read_text(encoding="utf-8"))
    assert manifest["game"]=="Tiny Heist"
    concepts={a["concept"] for a in manifest["assets"]}
    assert {"player_character","guard_enemy","security_camera","security_laser","door","floor_tile"} <= concepts

def test_out_of_scope_metrics_and_questions_do_not_become_assets(tmp_path: Path) -> None:
    gdd=tmp_path/"tiny.md"
    gdd.write_text(
        "# Tiny Heist\n\n## Art Direction\nStylized 3D with readable silhouettes and cyan lighting.\n\n"
        "## High Concept\nThe player avoids guards, cameras, lasers and doors.\n\n"
        "## Explicitly out of MVP\nCharacter equipment inventory.\nComplex guard alert networks.\n\n"
        "## Prototype Success Metrics\nPlayers can predict guard/camera movement after a short introduction.\n\n"
        "## Decisions Required Before / During the GDD\nHow are hints monetized without bypassing the puzzle?\n",
        encoding="utf-8",
    )
    out=tmp_path/"out"
    plan(gdd,out)
    manifest=yaml.safe_load((out/"ASSET_MANIFEST.yaml").read_text(encoding="utf-8"))
    concepts={a["concept"] for a in manifest["assets"]}
    assert concepts=={"player_character","guard_enemy","security_camera","security_laser","door"}

def test_unresolved_style_question_blocks_without_override(tmp_path: Path) -> None:
    gdd=tmp_path/"style.md"
    gdd.write_text(
        "# Tiny Heist\n\n## Art Direction\nClean top-down readability with strong silhouettes.\n\n"
        "## Decisions Required Before / During the GDD\nWhat tone should the heist use: noir, cartoon, or spy-tech?\n",
        encoding="utf-8",
    )
    assert plan(gdd,tmp_path/"blocked")["status"]=="blocked"
    style_file=tmp_path/"style.yaml"
    style_file.write_text(
        "schema_version: 1\n"
        "game:\n  title: Tiny Heist\n"
        "visual_style:\n  status: resolved\n  summary:\n    - premium stylized 3D playful suspense\n"
        "sources:\n  - type: user\n    reference: approved style\n"
        "confidence: explicit\n",
        encoding="utf-8",
    )
    assert plan(gdd,tmp_path/"resolved",style_file=style_file)["status"]=="awaiting_human_approval"

def test_existing_repository_asset_suppresses_duplicate_generation_and_sets_dimensions(tmp_path: Path) -> None:
    game=tmp_path/"game"
    assets=game/"Resources"
    assets.mkdir(parents=True)
    Image.new("RGBA",(512,256)).save(assets/"security_camera.png")
    gdd=tmp_path/"gdd.md"
    gdd.write_text(
        "# Repo Game\n\n## Art Direction\nStylized 3D with readable silhouettes.\n\n"
        "## Security\nSecurity cameras watch corridors.\n",
        encoding="utf-8",
    )
    out=tmp_path/"out"
    plan(gdd,out,game_root=game)
    manifest=yaml.safe_load((out/"ASSET_MANIFEST.yaml").read_text(encoding="utf-8"))
    camera=next(a for a in manifest["assets"] if a["concept"]=="security_camera")
    assert camera["generation"]["required"] is False
    assert camera["dimensions"]["source"]["type"]=="existing_resource"
    assert (camera["dimensions"]["width"],camera["dimensions"]["height"])==(512,256)


def test_build_outputs_are_excluded_and_missing_reference_becomes_target_asset(tmp_path: Path) -> None:
    game=tmp_path/"game"
    source=game/"Source"
    source.mkdir(parents=True)
    build=game/"build_x64"/"bin"
    build.mkdir(parents=True)
    Image.new("RGBA",(64,64)).save(build/"door.png")
    (source/"Scene.cpp").write_text('auto door = Sprite::create("Content/Art/Entities/door.png");',encoding="utf-8")

    gdd=tmp_path/"gdd.md"
    gdd.write_text(
        "# Repo Missing Game\n\n## Art Direction\nStylized 3D with readable top-down silhouettes.\n\n"
        "## High Concept\nThe player avoids guards and doors.\n",
        encoding="utf-8",
    )

    out=tmp_path/"out"
    plan(gdd,out,game_root=game)
    manifest=yaml.safe_load((out/"ASSET_MANIFEST.yaml").read_text(encoding="utf-8"))
    door=next(a for a in manifest["assets"] if a["concept"]=="door")
    assert door["generation"]["required"] is True
    assert door["generation"]["reason"]=="missing_repository_reference"
    assert door["target_path"]=="Content/Art/Entities/door.png"
    assert door["filename"]=="door.png"

    scan=yaml.safe_load((out/"STATE.yaml").read_text(encoding="utf-8"))
    assert scan["inputs"]["repository_scan"]["enabled"] is True
